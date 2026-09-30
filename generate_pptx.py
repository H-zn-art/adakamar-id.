import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # Set 16:9 widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6]

    # Brand Colors
    COLOR_PRIMARY = RGBColor(159, 60, 22)      # Terracotta #9f3c16
    COLOR_PRIMARY_DARK = RGBColor(122, 46, 16) # Deep Terracotta
    COLOR_DARK = RGBColor(26, 27, 34)          # Charcoal #1a1b22
    COLOR_MUTED = RGBColor(120, 113, 108)      # Text Muted #78716c
    COLOR_LIGHT_BG = RGBColor(248, 246, 244)   # Warm Cream #f8f6f4
    COLOR_WHITE = RGBColor(255, 255, 255)
    COLOR_CARD_BORDER = RGBColor(228, 224, 220)
    COLOR_SOFT_ORANGE = RGBColor(255, 219, 207) # #ffdbcf
    COLOR_EMERALD = RGBColor(22, 163, 74)

    def add_header(slide, badge_text, title_text, subtitle_text):
        # Header Badge
        badge_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.4))
        tf_b = badge_box.text_frame
        tf_b.word_wrap = True
        tf_b.margin_left = tf_b.margin_top = tf_b.margin_right = tf_b.margin_bottom = 0
        p_b = tf_b.paragraphs[0]
        p_b.text = f"✦ {badge_text.upper()}"
        p_b.font.name = "Arial"
        p_b.font.size = Pt(11)
        p_b.font.bold = True
        p_b.font.color.rgb = COLOR_PRIMARY

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.7))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.name = "Arial"
        p_t.font.size = Pt(24)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_DARK

        # Subtitle
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(0.4))
        tf_s = sub_box.text_frame
        tf_s.word_wrap = True
        tf_s.margin_left = tf_s.margin_top = tf_s.margin_right = tf_s.margin_bottom = 0
        p_s = tf_s.paragraphs[0]
        p_s.text = subtitle_text
        p_s.font.name = "Arial"
        p_s.font.size = Pt(12)
        p_s.font.color.rgb = COLOR_MUTED

    def add_card(slide, left, top, width, height, title, body_paragraphs, is_primary=False):
        # Card Background Shape
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.adjustments[0] = 0.06
        if is_primary:
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor(255, 245, 242)
            shape.line.color.rgb = COLOR_PRIMARY
            shape.line.width = Pt(1.5)
        else:
            shape.fill.solid()
            shape.fill.fore_color.rgb = COLOR_LIGHT_BG
            shape.line.color.rgb = COLOR_CARD_BORDER
            shape.line.width = Pt(1)

        # Card Content Text
        pad = Inches(0.25)
        tb = slide.shapes.add_textbox(left + pad, top + pad, width - (pad * 2), height - (pad * 2))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_title = tf.paragraphs[0]
        p_title.text = title
        p_title.font.name = "Arial"
        p_title.font.size = Pt(14)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_PRIMARY if is_primary else COLOR_DARK

        for text in body_paragraphs:
            p = tf.add_paragraph()
            p.text = text
            p.font.name = "Arial"
            p.font.size = Pt(11)
            p.font.color.rgb = COLOR_DARK
            p.space_before = Pt(6)

    def add_footer(slide, current_num, total_num=12):
        fb = slide.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(11.733), Inches(0.4))
        tf = fb.text_frame
        p = tf.paragraphs[0]
        p.text = f"adakamar.id — Platform Kurasi Homestay Autentik Yogyakarta  |  Slide {current_num:02d} / {total_num:02d}"
        p.font.name = "Arial"
        p.font.size = Pt(10)
        p.font.color.rgb = COLOR_MUTED

    # ══════════════════════════════════════════════════
    # SLIDE 1: COVER SLIDE
    # ══════════════════════════════════════════════════
    s1 = prs.slides.add_slide(blank_slide_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = COLOR_DARK
    bg1.line.fill.background()

    tb1 = s1.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(11.3), Inches(4.5))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "✦ LAPORAN AKHIR MAGANG & SHOWCASE SISTEM"
    p.font.name = "Arial"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_SOFT_ORANGE

    p = tf1.add_paragraph()
    p.text = "adakamar.id"
    p.font.name = "Arial"
    p.font.size = Pt(48)
    p.font.bold = True
    p.font.color.rgb = COLOR_WHITE
    p.space_before = Pt(10)

    p = tf1.add_paragraph()
    p.text = "Platform Kurasi Homestay Autentik & Ekosistem Pariwisata Yogyakarta"
    p.font.name = "Arial"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = COLOR_SOFT_ORANGE
    p.space_before = Pt(6)

    p = tf1.add_paragraph()
    p.text = "Solusi reservasi terkurasi untuk homestay tradisional Jawa (Joglo, Limasan, Villa Etnik) dengan integrasi WhatsApp Concierge, kupon promo dinamis, dan meja redaksi budaya lokal."
    p.font.name = "Arial"
    p.font.size = Pt(13)
    p.font.color.rgb = RGBColor(214, 211, 209)
    p.space_before = Pt(16)

    p = tf1.add_paragraph()
    p.text = "Tech Stack: Next.js 16 (App Router) • React 19 • TypeScript • NestJS • Prisma ORM • MySQL"
    p.font.name = "Arial"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = RGBColor(254, 215, 170)
    p.space_before = Pt(20)

    add_footer(s1, 1)

    # ══════════════════════════════════════════════════
    # SLIDE 2: LATAR BELAKANG & PERMASALAHAN
    # ══════════════════════════════════════════════════
    s2 = prs.slides.add_slide(blank_slide_layout)
    add_header(s2, "Latar Belakang & Permasalahan", "Tantangan Ekosistem Homestay di Yogyakarta", "Tiga masalah utama yang dihadapi wisatawan dan pemilik homestay lokal saat ini")
    
    add_card(s2, Inches(0.8), Inches(2.2), Inches(3.6), Inches(4.3),
             "1. Dominasi OTA Homogen",
             ["• Platform OTA global (Agoda, Traveloka) dipenuhi hotel jaringan steril yang seragam.",
              "• Keunikan budaya omah joglo dan suasana pedesaan Mataram tenggelam dan sulit ditemukan.",
              "• Potongan komisi tinggi (15% - 25%) membebani pemilik usaha kecil lokal."])
              
    add_card(s2, Inches(4.8), Inches(2.2), Inches(3.6), Inches(4.3),
             "2. Kesenjangan Digitalisasi Host",
             ["• Pemilik rumah bersejarah di Bantul, Sleman, & Kulon Progo minim kapabilitas digital.",
              "• Sulit mengelola ketersediaan kamar dan promosi online mandiri.",
              "• Tamu kesulitan mendapatkan konfirmasi ketersediaan kamar yang pasti & cepat."])

    add_card(s2, Inches(8.8), Inches(2.2), Inches(3.6), Inches(4.3),
             "3. Ekspektasi Tamu Modern",
             ["• Wisatawan mencari 'filosofi ngaso' (ketenangan jiwa) tapi ragu akan standar kebersihan.",
              "• Butuh jaminan kurasi fisik (bebas zonk), sanitasi higienis, AC, air panas, dan Wi-Fi.",
              "• Menginginkan alur pemesanan langsung tanpa registrasi akun yang berbelit."])
    add_footer(s2, 2)

    # ══════════════════════════════════════════════════
    # SLIDE 3: VALUE PROPOSITION & VISI
    # ══════════════════════════════════════════════════
    s3 = prs.slides.add_slide(blank_slide_layout)
    add_header(s3, "Solusi & Nilai Unggul", "Visi & Nilai Unik adakamar.id", "Menghubungkan pelancong bernyawa seni dengan kehangatan keramahan Mataram")

    add_card(s3, Inches(0.8), Inches(2.2), Inches(5.6), Inches(4.3),
             "🏛️ Kurasi Fisik 100% Autentik",
             ["• Menjamin keaslian arsitektur: sokoguru kayu jati asli, atap joglo/limasan, dan ubin tegel motif klasik.",
              "• Standar kenyamanan butik modern: kasur king-coil higienis, pendingin udara, air panas bertekanan.",
              "• Dokumentasi visual nyata: galeri foto beresolusi tinggi tanpa sudut distorsi atau rekayasa kamera.",
              "• Melestarikan warisan cagar budaya melalui pemanfaatan pariwisata berkelanjutan."],
             is_primary=True)

    add_card(s3, Inches(6.8), Inches(2.2), Inches(5.6), Inches(4.3),
             "🤝 Kemitraan Warga & Direct WhatsApp",
             ["• Reservasi instan terhubung langsung ke pengelola via WhatsApp concierge tanpa biaya tersembunyi.",
              "• Format pesan otomatis terstruktur dengan rincian tanggal, jumlah tamu, dan potongan kupon hemat.",
              "• Lebih dari 90% perputaran ekonomi mengalir langsung ke keluarga pemilik dan warga desa setempat.",
              "• Komisi terjangkau dan transparan untuk mendukung pertumbuhan UMKM pariwisata Jogja."],
             is_primary=True)
    add_footer(s3, 3)

    # ══════════════════════════════════════════════════
    # SLIDE 4: USER PERSONAS & ROLE RBAC
    # ══════════════════════════════════════════════════
    s4 = prs.slides.add_slide(blank_slide_layout)
    add_header(s4, "Aktor & Hak Akses", "Struktur Role & Alur Pengguna (RBAC)", "Pemisahan hak akses terisolasi untuk keamanan dan efisiensi operasional")

    add_card(s4, Inches(0.8), Inches(2.2), Inches(3.6), Inches(4.3),
             "👤 1. Pengunjung (Public)",
             ["• Wisatawan yang mencari homestay autentik & panduan wisata.",
              "• Filter multi-parameter (kawasan, harga, kapasitas, fasilitas).",
              "• Penggunaan kupon diskon dinamis pada form booking.",
              "• Kirim formulir reservasi langsung ke WhatsApp host tanpa wajib buat akun login."])

    add_card(s4, Inches(4.8), Inches(2.2), Inches(3.6), Inches(4.3),
             "✍️ 2. Penulis (PENULIS)",
             ["• Kurator konten budaya & artikel panduan wisata Yogyakarta.",
              "• Meja redaksi terisolasi: buat naskah, simpan draf, & ajukan review.",
              "• Unggah foto sampul langsung dari komputer.",
              "• Pasang & buat tagar artikel baru secara mandiri tanpa terblokir hak akses."])

    add_card(s4, Inches(8.8), Inches(2.2), Inches(3.6), Inches(4.3),
             "🛡️ 3. Administrator (ADMIN)",
             ["• Pengelola utama operasional, editorial, & kampanye bisnis.",
              "• Manajemen katalog homestay, fasilitas, & 6 wilayah wisata.",
              "• Tinjau, setujui, atau minta revisi naskah dari Penulis.",
              "• Manajemen kupon promo, kuota klaim, & modal box inquiry reservasi."])
    add_footer(s4, 4)

    # ══════════════════════════════════════════════════
    # SLIDE 5: TECH STACK & ARSITEKTUR
    # ══════════════════════════════════════════════════
    s5 = prs.slides.add_slide(blank_slide_layout)
    add_header(s5, "Arsitektur & Teknologi", "Full-Stack Architecture & Modern Tech Stack", "Pemisahan bersih Client-Server untuk performa kilat, SEO prima, dan keamanan tingkat tinggi")

    add_card(s5, Inches(0.8), Inches(2.2), Inches(5.6), Inches(4.3),
             "Frontend: Next.js 16 + React 19",
             ["• Next.js 16 (App Router): SSR & Static Rendering untuk optimasi SEO nomor 1 di Google.",
              "• React 19: Server Actions, Suspense boundaries, dan concurrent transitions.",
              "• TailwindCSS v4: Styling performan tinggi dengan sistem token warna Terracotta.",
              "• Lucide React: Iconography modern yang konsisten dan informatif.",
              "• Local File Upload Route (/api/upload): Upload multipart foto langsung ke public/uploads/."],
             is_primary=False)

    add_card(s5, Inches(6.8), Inches(2.2), Inches(5.6), Inches(4.3),
             "Backend: NestJS + Prisma ORM + MySQL",
             ["• NestJS 10: Arsitektur enterprise modular (Controllers, Services, DTOs, Guards).",
              "• Prisma ORM: Type-safe database queries, zero risk of SQL Injection.",
              "• MySQL 8.0: Basis data relasional terindeks untuk konsistensi data reservasi.",
              "• JWT + Bcrypt: Otentikasi token aman dengan Role-Based Access Guards (Admin & Penulis).",
              "• RESTful API: Endpoint terstandarisasi untuk homestay, promo, artikel, & inquiry."],
             is_primary=False)
    add_footer(s5, 5)

    # ══════════════════════════════════════════════════
    # SLIDE 6: MODUL 1 - HOMESTAY MARKETPLACE
    # ══════════════════════════════════════════════════
    s6 = prs.slides.add_slide(blank_slide_layout)
    add_header(s6, "Modul Fitur Utama 1", "Katalog & Reservasi Homestay (Marketplace)", "Fitur eksplorasi penginapan ramah pengguna dengan alur transaksi transparan")

    add_card(s6, Inches(0.8), Inches(2.2), Inches(3.6), Inches(4.3),
             "🔍 Filter Cerdas 6 Kawasan",
             ["• Eksplorasi 6 kawasan wisata Jogja: Malioboro, Prawirotaman, Kaliurang, Kotagede, Sleman, Bantul.",
              "• Filter rentang tarif per malam (budget s/d luxury).",
              "• Filter kapasitas kamar tidur, kamar mandi, dan jumlah tamu rombongan.",
              "• Filter fasilitas (kolam renang, dapur, Wi-Fi, AC, sarapan tradisional)."])

    add_card(s6, Inches(4.8), Inches(2.2), Inches(3.6), Inches(4.3),
             "📸 Halaman Detail Interaktif",
             ["• Galeri foto resolusi tinggi dengan visual arsitektur lengkap.",
              "• Aturan menginap (check-in/out, kebijakan keluarga, larangan merokok).",
              "• Ulasan dan rating bintang terverifikasi dari tamu terdahulu.",
              "• Peta lokasi interaktif (Leaflet Maps) dan titik koordinat akurat."])

    add_card(s6, Inches(8.8), Inches(2.2), Inches(3.6), Inches(4.3),
             "💬 Booking Concierge Widget",
             ["• Kalkulator otomatis durasi malam dan estimasi subtotal.",
              "• Pemasangan kode kupon promo dengan potongan harga instan.",
              "• Begitu form dikirim, data langsung tersimpan ke basis data MySQL.",
              "• Membuka aplikasi WhatsApp dengan pesan terformat rapi untuk konfirmasi host."])
    add_footer(s6, 6)

    # ══════════════════════════════════════════════════
    # SLIDE 7: MODUL 2 - PROMO & KUPON ENGINE
    # ══════════════════════════════════════════════════
    s7 = prs.slides.add_slide(blank_slide_layout)
    add_header(s7, "Modul Fitur Utama 2", "Sistem Promo & Kupon Diskon Dinamis", "Validasi aturan promosi otomatis dengan pelacakan pemakaian kuota real-time")

    add_card(s7, Inches(0.8), Inches(2.2), Inches(5.6), Inches(4.3),
             "Mekanisme Kupon Diskon",
             ["• Mendukung jenis diskon Persentase (%) atau Potongan Tetap (Rp).",
              "• Parameter syarat transaksi: Minimum Biaya Booking & Batas Maksimal Diskon.",
              "• Validasi tanggal berlaku (Start Date & End Date) diuji otomatis pada server.",
              "• Penargetan khusus: Kupon promo dapat dikaitkan khusus ke unit homestay tertentu (PRD Seksi 13).",
              "• Halaman promo publik (/promo) menampilkan kartu voucher dengan tombol Salin & Gunakan."],
             is_primary=False)

    add_card(s7, Inches(6.8), Inches(2.2), Inches(5.6), Inches(4.3),
             "Pelacakan Pemakaian Kuota (usedCount)",
             ["• Endpoint Baru: PATCH /api/promos/use/:code untuk mencatat klaim pemakaian kupon.",
              "• Atomic Database Increment: Menambah usedCount di MySQL setiap kali kupon diterapkan.",
              "• Sinkronisasi Dua Arah: Tamu mengklik 'Gunakan' di form booking langsung memicu usePromo dan update kuota.",
              "• Admin Monitoring: Dashboard promo menyajikan 'Total Voucher Digunakan' dan visual progress bar pemakaian kuota."],
             is_primary=True)
    add_footer(s7, 7)

    # ══════════════════════════════════════════════════
    # SLIDE 8: MODUL 3 - EDITORIAL CMS PENULIS
    # ══════════════════════════════════════════════════
    s8 = prs.slides.add_slide(blank_slide_layout)
    add_header(s8, "Modul Fitur Utama 3", "Meja Redaksi & Editorial Penulis (CMS Artikel)", "Penyusunan naskah kebudayaan dan panduan wisata dengan alur kerja modern")

    add_card(s8, Inches(0.8), Inches(2.2), Inches(3.6), Inches(4.3),
             "📝 Editor Naskah Rapi",
             ["• Layout dua kolom terstruktur: Naskah di sisi kiri, Pengaturan editorial di sisi kanan.",
              "• Toolbar formatting fungsional: Bold, Italic, Underline, Heading 1/2, List, Quote, Link, Tabel.",
              "• Indikator live: Penghitung jumlah kata otomatis dan estimasi durasi baca artikel pembaca.",
              "• Optimasi SEO: Input judul Google search dan meta deskripsi naskah."])

    add_card(s8, Inches(4.8), Inches(2.2), Inches(3.6), Inches(4.3),
             "🏷️ Sistem Tagar Dinamis",
             ["• Menampilkan daftar saran tagar populer dari sistem yang dapat diklik untuk pasang/lepas.",
              "• Penulis dapat membuat tagar kustom baru secara instan di form.",
              "• Hak akses backend diperbarui (UserRole.PENULIS diizinkan membuat tag tanpa error 403).",
              "• Sinkronisasi relasi tagId lengkap saat simpan draf maupun update naskah."])

    add_card(s8, Inches(8.8), Inches(2.2), Inches(3.6), Inches(4.3),
             "🖼️ Unggah Foto dari Komputer",
             ["• Penulis dapat mengunggah foto sampul (thumbnail) langsung dari laptop/komputer via drag-and-drop.",
              "• Terintegrasi ke /api/upload lokal — meniadakan kebutuhan mencari link URL gambar luar.",
              "• Pratinjau thumbnail instan dengan tombol 'Ganti Foto' yang fleksibel.",
              "• Pilihan tab URL eksternal tetap tersedia sebagai opsi cadangan."])
    add_footer(s8, 8)

    # ══════════════════════════════════════════════════
    # SLIDE 9: MODUL 4 - ADMIN COMMAND CENTER
    # ══════════════════════════════════════════════════
    s9 = prs.slides.add_slide(blank_slide_layout)
    add_header(s9, "Modul Fitur Utama 4", "Admin Command Center & Pustaka Media", "Pusat kendali operasional, galeri aset visual lokal, dan penanganan reservasi tamu")

    add_card(s9, Inches(0.8), Inches(2.2), Inches(5.6), Inches(4.3),
             "Pustaka Media (Impor File Manual)",
             ["• Tombol Header: + Impor Foto Manual di halaman Pustaka Media.",
              "• Unggah berkas gambar langsung dari perangkat komputer lokal (JPG, PNG, WebP, GIF).",
              "• Dukungan multi-file: Memilih dan mengunggah beberapa foto sekaligus dalam satu langkah.",
              "• Pratinjau thumbnail, informasi ukuran file (KB/MB), dan opsi pembatalan sebelum disimpan.",
              "• Foto otomatis tersimpan ke /public/uploads/ dan ditautkan ke penginapan tujuan di database."],
             is_primary=True)

    add_card(s9, Inches(6.8), Inches(2.2), Inches(5.6), Inches(4.3),
             "Modal Box Reservasi Tamu",
             ["• Penggantian lembar drawer samping menjadi Modal Box Dialog Terpusat yang bersih dan elegan.",
              "• Rincian Lengkap: Kode reservasi, nama tamu, kontak WhatsApp, email, unit homestay, tanggal, durasi, total biaya.",
              "• Informasi Kupon Promo: Menampilkan potongan voucher yang dipakai tamu pada pemesanan.",
              "• Aksi Cepat Status: Tombol instan untuk konfirmasi, selesaikan, atau batalkan pemesanan.",
              "• Tombol 'Chat Tamu via WhatsApp' langsung membuka nomor telepon tamu dengan 1 klik."],
             is_primary=True)
    add_footer(s9, 9)

    # ══════════════════════════════════════════════════
    # SLIDE 10: SKEMA BASIS DATA & RELASI
    # ══════════════════════════════════════════════════
    s10 = prs.slides.add_slide(blank_slide_layout)
    add_header(s10, "Struktur Basis Data", "Skema Relasional Komprehensif (Prisma ORM)", "Pemodelan data terstandarisasi untuk integritas relasi dan performa query optimal")

    add_card(s10, Inches(0.8), Inches(2.2), Inches(2.7), Inches(4.3),
             "Entitas Pengguna",
             ["• User: id, name, email, password, role (ADMIN/PENULIS), bio, phone, avatarUrl.",
              "• Relasi 1-to-many ke Article untuk tracking kepemilikan naskah penulis."])

    add_card(s10, Inches(3.8), Inches(2.2), Inches(2.7), Inches(4.3),
             "Entitas Homestay",
             ["• Property: id, name, slug, price, address, capacity, rules, status.",
              "• PropertyImage (1-to-many).",
              "• Facility & PropertyFacility (many-to-many).",
              "• LocationArea & Category."])

    add_card(s10, Inches(6.8), Inches(2.2), Inches(2.7), Inches(4.3),
             "Entitas Promo",
             ["• Promo: code, title, discountPercent, discountAmount, quota, usedCount, dates.",
              "• PromoProperty: Pivot table many-to-many penginapan terkait kupon."])

    add_card(s10, Inches(9.8), Inches(2.2), Inches(2.7), Inches(4.3),
             "Artikel & Inquiry",
             ["• Article: title, slug, content, excerpt, status, views, publishedAt.",
              "• ArticleCategory & Tag.",
              "• ArticleTagPivot.",
              "• Inquiry: guest data, dates, notes, status."])
    add_footer(s10, 10)

    # ══════════════════════════════════════════════════
    # SLIDE 11: REVISI BERANDA: TENTANG KAMI
    # ══════════════════════════════════════════════════
    s11 = prs.slides.add_slide(blank_slide_layout)
    add_header(s11, "Penyempurnaan Antarmuka", "Revisi Beranda: Dari Host Menjadi 'Tentang Kami'", "Menyelaraskan pesan halaman utama dengan filosofi kurasi dan kenyamanan wisatawan")

    add_card(s11, Inches(0.8), Inches(2.2), Inches(5.6), Inches(4.3),
             "Pembaruan Navigasi Utama (Navbar)",
             ["• Item navigasi ke-6 yang sebelumnya berlabel 'Tuan Rumah' (/buka-homestay) kini diganti menjadi 'Tentang Kami'.",
              "• Mengarahkan ke rute fungsional aktif: /tentang-kami.",
              "• Menggunakan ikon informatif terstandarisasi (Lucide Info icon).",
              "• Footer website turut menyertakan tautan langsung ke halaman Tentang Kami.",
              "• Halaman /tentang-kami telah diuji dan menghasilkan HTTP Status 200 OK."],
             is_primary=False)

    add_card(s11, Inches(6.8), Inches(2.2), Inches(5.6), Inches(4.3),
             "Banner Narasi Filosofis 'Ngaso'",
             ["• Mengubah banner bawah beranda menjadi Kisah & Kurasi Autentik adakamar.id.",
              "• Menonjolkan 3 pilar kurasi: 100% Homestay Terkurasi, Pemberdayaan Warga Desa, Standar Butik Higienis.",
              "• Menghadirkan ajakan aksi yang jelas: Tombol 'Pelajari Kisah Kami' dan 'Jelajahi Homestay'.",
              "• Mengedepankan narasi budaya restorasi omah limasan & joglo lawasan khas Daerah Istimewa Yogyakarta."],
             is_primary=True)
    add_footer(s11, 11)

    # ══════════════════════════════════════════════════
    # SLIDE 12: KESIMPULAN & ROADMAP PENGEMBANGAN
    # ══════════════════════════════════════════════════
    s12 = prs.slides.add_slide(blank_slide_layout)
    bg12 = s12.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg12.fill.solid()
    bg12.fill.fore_color.rgb = COLOR_DARK
    bg12.line.fill.background()

    add_header(s12, "Penutup & Masa Depan", "Kesimpulan & Roadmap Pengembangan Sistem", "Pencapaian proyek magang dan rencana perluasan platform adakamar.id ke depan")
    
    # White text adjustments for dark slide header
    for shape in s12.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                if p.text.startswith("Tantangan") or "Kesimpulan" in p.text:
                    p.font.color.rgb = COLOR_WHITE
                elif "Pencapaian" in p.text:
                    p.font.color.rgb = RGBColor(214, 211, 209)

    add_card(s12, Inches(0.8), Inches(2.2), Inches(3.6), Inches(4.3),
             "💳 Payment Gateway QRIS",
             ["• Integrasi gerbang pembayaran otomatis (Midtrans / Xendit).",
              "• Menyediakan opsi pembayaran instan via QRIS, Virtual Account, & E-Wallet bagi tamu yang memilih bayar lunas di awal.",
              "• Notifikasi otomatis status pembayaran sukses ke sistem."])

    add_card(s12, Inches(4.8), Inches(2.2), Inches(3.6), Inches(4.3),
             "📅 Kalender iCal Sync",
             ["• Sinkronisasi ketersediaan kamar dua arah dengan Airbnb dan Google Calendar.",
              "• Mencegah double booking secara otomatis saat homestay dipesan lewat kanal berbeda.",
              "• Pembaruan jadwal sewa secara otomatis setiap 15 menit."])

    add_card(s12, Inches(8.8), Inches(2.2), Inches(3.6), Inches(4.3),
             "📱 Host Mobile PWA",
             ["• Pengembangan aplikasi Progressive Web App khusus pemilik homestay.",
              "• Notifikasi push seketika saat ada inquiry tamu baru masuk.",
              "• Kemudahan mengubah tarif musiman (peak season) langsung dari smartphone host."])
    add_footer(s12, 12)

    # Save to disk
    output_filename = "Presentasi_Project_adakamar_id.pptx"
    output_path = os.path.join(os.getcwd(), output_filename)
    prs.save(output_path)
    print(f"Presentation successfully saved to: {output_path}")

if __name__ == "__main__":
    create_presentation()
